import os, uuid
from pathlib import Path
from flask import Flask, render_template, request, send_file, flash, redirect, url_for
from werkzeug.utils import secure_filename
import fitz
import pdfplumber
from docx import Document
from openpyxl import Workbook, load_workbook
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm

ROOT = Path(__file__).resolve().parent
UPLOADS = Path(os.getenv("UPLOAD_DIR", ROOT / "uploads"))
OUTPUTS = Path(os.getenv("OUTPUT_DIR", ROOT / "outputs"))
UPLOADS.mkdir(exist_ok=True, parents=True)
OUTPUTS.mkdir(exist_ok=True, parents=True)
MAX_MB = int(os.getenv("MAX_UPLOAD_MB", "20"))
app = Flask(__name__)
app.secret_key = os.getenv("SECRET_KEY", "change-this-secret-before-production")
app.config["MAX_CONTENT_LENGTH"] = MAX_MB * 1024 * 1024
ALLOWED = {".pdf", ".docx", ".xlsx", ".xlsm"}

def save_file(file):
    if not file or not file.filename: raise ValueError("Please choose a file.")
    name = secure_filename(file.filename)
    ext = Path(name).suffix.lower()
    if ext not in ALLOWED: raise ValueError("Supported: PDF, DOCX, XLSX and XLSM.")
    path = UPLOADS / (uuid.uuid4().hex + "_" + name)
    file.save(path)
    return path

def out(ext): return OUTPUTS / (uuid.uuid4().hex + ext)

def pdf_word(src, dest):
    from pdf2docx import Converter
    cv = Converter(str(src))
    try: cv.convert(str(dest))
    finally: cv.close()

def pdf_excel(src, dest):
    wb = Workbook()
    ws = wb.active
    ws.title = "Extracted Data"
    row = 1
    with pdfplumber.open(src) as pdf:
        for page_no, page in enumerate(pdf.pages, 1):
            ws.cell(row, 1, f"Page {page_no}"); row += 1
            tables = page.extract_tables() or []
            for table in tables:
                for vals in table:
                    for col, val in enumerate(vals or [], 1):
                        ws.cell(row, col, val.strip() if isinstance(val, str) else val)
                    row += 1
                row += 1
            if not tables:
                for line in (page.extract_text() or "").splitlines():
                    ws.cell(row, 1, line); row += 1
            row += 1
    wb.save(dest)

def word_pdf(src, dest):
    doc = Document(str(src)); styles = getSampleStyleSheet(); story = []
    for p in doc.paragraphs:
        if p.text.strip():
            txt = p.text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            story += [Paragraph(txt, styles["BodyText"]), Spacer(1, 3*mm)]
    for tbl in doc.tables:
        data = [[c.text for c in r.cells] for r in tbl.rows]
        if data:
            t = Table(data, repeatRows=1)
            t.setStyle(TableStyle([("GRID",(0,0),(-1,-1),.4,colors.grey),("BACKGROUND",(0,0),(-1,0),colors.lightgrey),("FONTSIZE",(0,0),(-1,-1),8)]))
            story += [t, Spacer(1, 5*mm)]
    if not story: story = [Paragraph("No extractable content.", styles["BodyText"])]
    SimpleDocTemplate(str(dest), pagesize=A4).build(story)

def excel_pdf(src, dest):
    wb = load_workbook(src, data_only=True, read_only=True); styles = getSampleStyleSheet(); story = []
    for ws in wb.worksheets:
        story.append(Paragraph(ws.title, styles["Heading2"]))
        data = [[("" if v is None else str(v)) for v in row] for row in ws.iter_rows(values_only=True)]
        data = [r for r in data if any(r)]
        if data:
            cols = min(max(map(len, data)), 8)
            data = [(r + [""]*cols)[:cols] for r in data]
            t = Table(data, repeatRows=1)
            t.setStyle(TableStyle([("GRID",(0,0),(-1,-1),.3,colors.grey),("BACKGROUND",(0,0),(-1,0),colors.lightgrey),("FONTSIZE",(0,0),(-1,-1),7)]))
            story += [t, Spacer(1, 6*mm)]
    wb.close()
    if not story: story = [Paragraph("No spreadsheet data.", styles["BodyText"])]
    SimpleDocTemplate(str(dest), pagesize=landscape(A4), leftMargin=8*mm, rightMargin=8*mm).build(story)

@app.get("/")
def index(): return render_template("index.html")

@app.post("/convert")
def convert():
    src = None
    try:
        src = save_file(request.files.get("file")); ext = src.suffix.lower(); action = request.form.get("action")
        if action == "pdf_word" and ext == ".pdf": dest = out(".docx"); pdf_word(src, dest)
        elif action == "pdf_excel" and ext == ".pdf": dest = out(".xlsx"); pdf_excel(src, dest)
        elif action == "word_pdf" and ext == ".docx": dest = out(".pdf"); word_pdf(src, dest)
        elif action == "excel_pdf" and ext in {".xlsx", ".xlsm"}: dest = out(".pdf"); excel_pdf(src, dest)
        else: raise ValueError("Selected tool does not match your file type.")
        return send_file(dest, as_attachment=True, download_name="converted" + dest.suffix)
    except Exception as e:
        flash("Conversion failed: " + str(e)); return redirect(url_for("index"))
    finally:
        if src: src.unlink(missing_ok=True)

@app.post("/pdf-tools")
def pdf_tools():
    src = None
    try:
        src = save_file(request.files.get("file"))
        if src.suffix.lower() != ".pdf": raise ValueError("Upload a PDF file.")
        action = request.form.get("action"); doc = fitz.open(src)
        if action == "rotate":
            angle = int(request.form.get("angle", "90"))
            if angle not in (90,180,270): raise ValueError("Rotation must be 90, 180 or 270.")
            for page in doc: page.set_rotation((page.rotation + angle) % 360)
        elif action == "split":
            indexes = []
            for part in request.form.get("pages","").split(","):
                part = part.strip()
                if not part: continue
                if "-" in part:
                    a,b = map(int,part.split("-",1))
                    if a < 1 or b < a or b > len(doc): raise ValueError("Invalid page range.")
                    indexes.extend(range(a-1,b))
                else:
                    p = int(part)
                    if p < 1 or p > len(doc): raise ValueError("Page number out of range.")
                    indexes.append(p-1)
            if not indexes: raise ValueError("Enter pages, e.g. 1-3,5.")
            new = fitz.open()
            for i in indexes: new.insert_pdf(doc, from_page=i, to_page=i)
            doc.close(); doc = new
        elif action == "merge":
            extra = save_file(request.files.get("extra_file"))
            try:
                if extra.suffix.lower() != ".pdf": raise ValueError("Second file must be a PDF.")
                with fitz.open(extra) as other: doc.insert_pdf(other)
            finally: extra.unlink(missing_ok=True)
        elif action == "add_text":
            page_no = int(request.form.get("page","1")); text = request.form.get("text","").strip()
            if not text or not 1 <= page_no <= len(doc): raise ValueError("Enter valid page and text.")
            doc[page_no-1].insert_text((72,72), text, fontsize=float(request.form.get("font_size","12")), color=(.1,.25,.7))
        elif action == "replace_text":
            page_no = int(request.form.get("page","1")); old = request.form.get("old_text","").strip(); new = request.form.get("new_text","")
            if not old or not 1 <= page_no <= len(doc): raise ValueError("Enter a valid page and exact existing text.")
            page = doc[page_no-1]; rects = page.search_for(old)
            if not rects: raise ValueError("Text not found; it may be scanned or split into pieces.")
            for rect in rects: page.add_redact_annot(rect, fill=(1,1,1))
            page.apply_redactions()
            for rect in rects: page.insert_text((rect.x0, rect.y1-2), new, fontsize=float(request.form.get("font_size","11")), color=(0,0,0))
        else: raise ValueError("Choose a PDF operation.")
        dest = out(".pdf"); doc.save(dest, garbage=4, deflate=True); doc.close()
        return send_file(dest, as_attachment=True, download_name="edited.pdf")
    except Exception as e:
        flash("PDF operation failed: " + str(e)); return redirect(url_for("index"))
    finally:
        if src: src.unlink(missing_ok=True)

@app.errorhandler(413)
def too_large(_):
    flash(f"File too large. Maximum size is {MAX_MB} MB.")
    return redirect(url_for("index"))

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.getenv("PORT","5000")))
